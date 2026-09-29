"""CSV ↔ Excel 互转。"""
import csv
from openpyxl import Workbook, load_workbook

# CSV → Excel
with open("data.csv", "w", newline="", encoding="utf-8") as f:
    csv.writer(f).writerows([["城市", "人口"], ["上海", 2487], ["北京", 2189], ["深圳", 1756]])

wb = Workbook()
ws = wb.active
with open("data.csv", encoding="utf-8") as f:
    for row in csv.reader(f):
        ws.append(row)
wb.save("城市.xlsx")

# Excel → CSV
rows = load_workbook("城市.xlsx").active.iter_rows(values_only=True)
with open("还原.csv", "w", newline="", encoding="utf-8") as f:
    csv.writer(f).writerows(rows)
print("互转完成，内容一致:", open("data.csv").read() == open("还原.csv").read())
