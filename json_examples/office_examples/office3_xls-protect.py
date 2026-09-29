"""工作表保护：锁定公式，开放输入列。"""
from openpyxl import Workbook
from openpyxl.styles import Protection

wb = Workbook()
ws = wb.active
ws.append(["数量", "单价", "金额"])
for i in range(1, 4):
    ws.append([None, 25.0, None])
    ws.cell(row=i + 1, column=3, value=f"=A{i + 1}*B{i + 1}")
for row in ws.iter_rows(min_row=2, min_col=1, max_col=2):  # 输入列解锁
    for cell in row:
        cell.protection = Protection(locked=False)
ws.protection.sheet = True
wb.save("受保护表.xlsx")
print("公式列已锁定，数量/单价列可编辑")
