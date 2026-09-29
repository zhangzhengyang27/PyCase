"""Word 表格：数据报表插入。"""
from docx import Document

doc = Document()
doc.add_heading("季度数据", level=1)
table = doc.add_table(rows=1, cols=3)
table.style = "Light Grid Accent 1"
for i, h in enumerate(["季度", "营收(万)", "同比"]):
    table.rows[0].cells[i].text = h
for row in [("Q1", 120, "+8%"), ("Q2", 135, "+12%"), ("Q3", 150, "+11%")]:
    cells = table.add_row().cells
    for i, v in enumerate(row):
        cells[i].text = str(v)
doc.save("季度数据.docx")
print("已生成含表格的 季度数据.docx")
