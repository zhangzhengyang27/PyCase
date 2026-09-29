"""Excel 样式：表头加粗白字 + 蓝底 + 定宽列。"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

wb = Workbook()
ws = wb.active
ws.append(["姓名", "部门", "工资"])
for r in [["张三", "研发", 25000], ["李四", "市场", 18000], ["王五", "研发", 27000]]:
    ws.append(r)
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill("solid", fgColor="4472C4")
for cell in ws[1]:
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center")
for col in "ABC":
    ws.column_dimensions[col].width = 14
wb.save("员工表.xlsx")
print("已生成带样式的 员工表.xlsx")
