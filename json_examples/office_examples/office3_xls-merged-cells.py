"""合并单元格：两级表头。"""
from openpyxl import Workbook
from openpyxl.styles import Alignment

wb = Workbook()
ws = wb.active
ws.merge_cells("A1:A2"); ws["A1"] = "月份"
ws.merge_cells("B1:C1"); ws["B1"] = "销售额"
ws["B2"], ws["C2"] = "线上", "线下"
for cell in ("A1", "B1"):
    ws[cell].alignment = Alignment(horizontal="center", vertical="center")
ws.append(["1月", 100, 60])
ws.append(["2月", 120, 75])
wb.save("分组表头.xlsx")
print("两级表头已生成")
