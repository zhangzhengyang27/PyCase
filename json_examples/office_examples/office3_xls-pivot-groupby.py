"""透视：按部门/月份双维聚合。"""
from collections import defaultdict
from openpyxl import Workbook

rows = [("研发", "1月", 120), ("研发", "2月", 135), ("市场", "1月", 90),
        ("市场", "2月", 110), ("研发", "1月", 60)]
pivot = defaultdict(float)
for dept, month, amount in rows:
    pivot[(dept, month)] += amount
depts = sorted({r[0] for r in rows})
months = sorted({r[1] for r in rows})

wb = Workbook()
ws = wb.active
ws.append(["部门"] + months)
for d in depts:
    ws.append([d] + [pivot.get((d, m), 0) for m in months])
wb.save("透视表.xlsx")
print("透视完成：", wb.sheetnames, ws.max_row - 1, "行")
