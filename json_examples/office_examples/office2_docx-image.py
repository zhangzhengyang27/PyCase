"""Word 插图：matplotlib 出图 → 插入 Word。"""
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches

fig, ax = plt.subplots(figsize=(5, 3))
ax.bar(["Q1", "Q2", "Q3"], [120, 135, 150], color="#4472C4")
fig.savefig("季度图.png", dpi=100)
plt.close(fig)

doc = Document()
doc.add_heading("季度营收", level=1)
doc.add_picture("季度图.png", width=Inches(4.5))
doc.save("季度报告.docx")
print("已生成含图的 季度报告.docx")
