"""Word 生成：会议纪要模板。"""
from docx import Document

doc = Document()
doc.add_heading("项目周会纪要", level=1)
doc.add_paragraph("时间：2026-09-23 10:00\n地点：3F 会议室")
doc.add_heading("一、本周进展", level=2)
doc.add_paragraph("爬虫模块联调完成", style="List Bullet")
doc.add_paragraph("数据看板上线", style="List Bullet")
doc.add_heading("二、风险项", level=2)
doc.add_paragraph("第三方接口限流，需申请配额", style="Intense Quote")
doc.save("周会纪要.docx")
print("已生成 周会纪要.docx")
