"""PPT 生成：三页周报幻灯片。"""
from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()
s1 = prs.slides.add_slide(prs.slide_layouts[0])
s1.shapes.title.text = "项目周报"
s1.placeholders[1].text = "2026-09-23 · 自动生成"

s2 = prs.slides.add_slide(prs.slide_layouts[1])
s2.shapes.title.text = "本周要点"
for line in ["爬虫模块联调完成", "示例库扩充至 3000+", "看板数据接入"]:
    s2.placeholders[1].text_frame.add_paragraph().text = line

s3 = prs.slides.add_slide(prs.slide_layouts[5])
s3.shapes.title.text = "数据概览"
s3.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2)).text_frame.text = "3120 个示例 · 13 个集合 · 测试全绿"
prs.save("项目周报.pptx")
print("已生成 项目周报.pptx（3 页）")
