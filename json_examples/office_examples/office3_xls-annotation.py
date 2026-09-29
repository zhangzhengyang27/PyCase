"""批注：给数据加口径说明。"""
from openpyxl import Workbook
from openpyxl.comments import Comment

wb = Workbook()
ws = wb.active
ws.append(["指标", "数值"])
ws["B2"] = 3120
ws["B2"].comment = Comment("示例库总量，2026-09-23 快照", "报表机器人")
ws["B3"] = 39
ws["B3"].comment = Comment("共享依赖包数（requirements.txt）", "报表机器人")
wb.save("带批注.xlsx")
print("批注已写入（悬停单元格可见）")
